import logging
try:
    from docx import Document
except ImportError:
    Document = None

class DocParser:
    def parse(self, content: bytes) -> str:
        """
        Extracts text from DOCX files using python-docx.
        """
        if Document is None:
            logging.warning("python-docx not installed, return empty string.")
            return ""

        try:
            # python-docx expects a file-like object
            import io
            f = io.BytesIO(content)
            doc = Document(f)
            
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            
            # Simple table extraction (flattened)
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells]
                    full_text.append(" | ".join(row_text))
            
            return "\n".join(full_text)
        except Exception as e:
            logging.error(f"Docx Parse error: {e}")
            return ""
